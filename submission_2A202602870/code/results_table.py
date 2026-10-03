"""results_table.py — lưu kết quả từng lần chạy ra JSON và điền experiments.xlsx từ mẫu."""
from __future__ import annotations

import json
import math
from pathlib import Path

FORMULA_COLS = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}


def save_result(result: dict, results_dir: str = "../results") -> str:
    d = Path(results_dir)
    d.mkdir(parents=True, exist_ok=True)
    p = d / f'{result["cfg"]["exp_id"]}.json'
    p.write_text(json.dumps({k: result[k] for k in ("cfg", "history", "summary")}, default=str, indent=1),
                 encoding="utf-8")
    return str(p)


def load_results(results_dir: str = "../results") -> list[dict]:
    rs = [json.loads(p.read_text(encoding="utf-8")) for p in Path(results_dir).glob("*.json")]
    return sorted(rs, key=lambda r: r["cfg"]["exp_id"])


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    row = {**result["cfg"], **result["summary"], "notes": notes}
    row["hidden"] = str(tuple(row["hidden"]))
    row["figure_file"] = f'figures/{row["exp_id"]}.png'
    if eval_scores:
        row["eval_acc"], row["eval_macro_f1"] = eval_scores["eval_acc"], eval_scores["eval_macro_f1"]
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    import openpyxl
    wb = openpyxl.load_workbook(template_path)    # không dùng data_only=True (mất công thức)
    ws = wb["Experiments"]
    cols = {c.value: c.column for c in ws[1] if c.value}
    for i, row in enumerate(rows, start=2):
        for k, v in row.items():
            if k in cols and k not in FORMULA_COLS:
                ws.cell(row=i, column=cols[k], value=None if isinstance(v, float) and math.isnan(v) else v)
    wb.save(out_path)
